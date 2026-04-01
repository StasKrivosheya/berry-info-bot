from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.knowledge_base.cli import main
from app.services.knowledge_base.markdown_render import infer_fallback_headers
from app.services.knowledge_base.parser import parse_knowledge_base


def _write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def _save_workbook(path: Path, workbook: Workbook) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def test_multiline_cells_render_and_manifest_is_generated(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    _write_text(
        input_dir / "faq.csv",
        'Question,Answer\n"How to order?","Step one.\n\nStep two."\n',
    )
    _write_text(
        config_path,
        """
[files."faq.csv"]
question_column_name = "Question"
answer_column_name = "Answer"
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 1
    assert result.failure_count == 0
    markdown_path = output_dir / "markdown" / "faq.md"
    manifest_path = output_dir / "manifest.json"
    assert markdown_path.exists()
    assert manifest_path.exists()

    markdown = markdown_path.read_text(encoding="utf-8")
    assert "# Faq" in markdown
    assert "### How to order?" in markdown
    assert "Step one.\n\nStep two." in markdown

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert "generated_at_utc" in manifest
    assert manifest["entries"][0]["source_file"] == "faq.csv"
    assert manifest["entries"][0]["source_format"] == "csv"
    assert manifest["entries"][0]["sheet_name"] is None
    assert manifest["entries"][0]["workbook_file"] is None
    assert manifest["entries"][0]["output_md_file"] == ["markdown/faq.md"]
    assert manifest["entries"][0]["content_hash_sha256"]


def test_manifest_hash_and_logical_id_are_stable_across_runs(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    _write_text(input_dir / "01-faq.csv", "Question,Answer\nQ1,A1\n")
    _write_text(
        config_path,
        """
[logical_id_overrides]
"01-faq" = "mapped-id"

[files."01-faq.csv"]
logical_id = "file-level-id"
question_column_name = "Question"
answer_column_name = "Answer"
""".strip(),
    )

    first = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )
    second = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert first.entries[0].logical_id == "file-level-id"
    assert second.entries[0].logical_id == "file-level-id"
    assert first.entries[0].output_md_file == ["markdown/01-faq.md"]
    assert second.entries[0].output_md_file == ["markdown/01-faq.md"]
    assert first.entries[0].content_hash_sha256 == second.entries[0].content_hash_sha256


def test_partial_failure_does_not_fail_batch_but_all_fail_returns_non_zero(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    _write_text(input_dir / "ok.csv", "Header,Value\nTopic,Text\n")
    _write_text(input_dir / "broken.csv", 'Question,Answer\n"Broken,Cell\n')
    _write_text(config_path, "")

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )
    assert result.success_count == 1
    assert result.failure_count == 1
    assert not result.all_files_failed

    exit_code_partial = main(
        [
            "--input-dir",
            str(input_dir),
            "--output-dir",
            str(output_dir),
            "--config",
            str(config_path),
        ]
    )
    assert exit_code_partial == 0

    only_broken_input = tmp_path / "raw_only_broken"
    only_broken_output = tmp_path / "processed_only_broken"
    _write_text(only_broken_input / "broken.csv", 'Question,Answer\n"Broken,Cell\n')
    exit_code_all_failed = main(
        [
            "--input-dir",
            str(only_broken_input),
            "--output-dir",
            str(only_broken_output),
            "--config",
            str(config_path),
        ]
    )
    assert exit_code_all_failed == 1


def test_parser_rerun_keeps_markdown_bytes_and_stable_manifest_entry_fields(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    _write_text(input_dir / "faq.csv", "Question,Answer\nQ1,A1\nQ2,A2\n")
    _write_text(
        config_path,
        """
[files."faq.csv"]
question_column_name = "Question"
answer_column_name = "Answer"
""".strip(),
    )

    parse_knowledge_base(input_dir=input_dir, output_dir=output_dir, config_path=config_path)
    first_markdown = (output_dir / "markdown" / "faq.md").read_bytes()
    first_manifest = json.loads((output_dir / "manifest.json").read_text(encoding="utf-8"))
    first_entry = first_manifest["entries"][0]

    parse_knowledge_base(input_dir=input_dir, output_dir=output_dir, config_path=config_path)
    second_markdown = (output_dir / "markdown" / "faq.md").read_bytes()
    second_manifest = json.loads((output_dir / "manifest.json").read_text(encoding="utf-8"))
    second_entry = second_manifest["entries"][0]

    assert first_markdown == second_markdown
    assert first_entry["source_file"] == second_entry["source_file"]
    assert first_entry["source_format"] == second_entry["source_format"]
    assert first_entry["sheet_name"] == second_entry["sheet_name"]
    assert first_entry["workbook_file"] == second_entry["workbook_file"]
    assert first_entry["output_md_file"] == second_entry["output_md_file"]
    assert first_entry["logical_id"] == second_entry["logical_id"]
    assert first_entry["category"] == second_entry["category"]
    assert first_entry["version"] == second_entry["version"]
    assert first_entry["row_count"] == second_entry["row_count"]
    assert first_entry["non_empty_cell_count"] == second_entry["non_empty_cell_count"]
    assert first_entry["content_hash_sha256"] == second_entry["content_hash_sha256"]


def test_xlsx_outline_sheet_renders_title_sections_and_label_value(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land Programs"
    sheet["A3"] = "Family Day"
    sheet["A3"].font = Font(bold=True)
    sheet["A4"] = "Spend the day together."
    sheet["A5"] = "- rides\n- farm"
    sheet["A7"] = "Price: 550 UAH"
    _save_workbook(input_dir / "program.xlsx", workbook)

    _write_text(
        config_path,
        """
[files."program.xlsx".sheets."Offers"]
parser_profile = "outline_sheet"
logical_id = "spring-program"
category = "offers"
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 1
    assert result.failure_count == 0
    markdown_path = output_dir / "markdown" / "program-offers.md"
    markdown = markdown_path.read_text(encoding="utf-8")
    assert markdown == (
        "# Berry Land Programs\n\n"
        "## Family Day\n\n"
        "Spend the day together.\n\n"
        "- rides\n"
        "- farm\n\n"
        "### Price\n\n"
        "550 UAH\n"
    )

    entry = result.entries[0]
    assert entry.source_file == "program.xlsx"
    assert entry.source_format == "xlsx"
    assert entry.sheet_name == "Offers"
    assert entry.workbook_file == "program.xlsx"
    assert entry.logical_id == "spring-program"


def test_xlsx_outline_sheet_supports_multiple_content_regions_and_skips_hidden_sheets(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalog"
    sheet.merge_cells("A1:J1")
    sheet["A1"] = "Berry Land Catalog"
    sheet["A3"] = "Left intro."
    sheet["H3"] = "Right intro."
    sheet["A5"] = "Schedule"
    sheet["A5"].font = Font(bold=True)
    sheet["H5"] = "Price: 550 UAH"
    hidden_sheet = workbook.create_sheet("Hidden")
    hidden_sheet.sheet_state = "hidden"
    hidden_sheet["A1"] = "Do not parse"
    _save_workbook(input_dir / "catalog.xlsx", workbook)

    _write_text(
        config_path,
        """
[files."catalog.xlsx".sheets."Catalog"]
parser_profile = "outline_sheet"
content_ranges = ["A:C", "H:J"]
logical_id = "catalog"
category = "sales"
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 1
    assert result.failure_count == 0
    markdown = (output_dir / "markdown" / "catalog-catalog.md").read_text(encoding="utf-8")
    assert markdown == (
        "# Berry Land Catalog\n\n"
        "Left intro.\n\n"
        "Right intro.\n\n"
        "## Schedule\n\n"
        "### Price\n\n"
        "550 UAH\n"
    )


def test_xlsx_outline_sheet_reports_ambiguous_short_rows(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ambiguous"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land"
    sheet["A2"] = "What is included?"
    sheet["A3"] = "- rides\n- farm"
    _save_workbook(input_dir / "ambiguous.xlsx", workbook)

    _write_text(
        config_path,
        """
[files."ambiguous.xlsx".sheets."Ambiguous"]
parser_profile = "outline_sheet"
logical_id = "ambiguous"
category = "sales"
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 0
    assert result.failure_count == 1
    assert result.errors[0].source_ref == "ambiguous.xlsx#Ambiguous"
    assert result.errors[0].diagnostics[0].code == "ambiguous_short_row"
    assert result.errors[0].diagnostics[0].row_numbers == (2,)


def test_xlsx_outline_sheet_reports_inline_header_blocks(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land"
    sheet["A2"] = "Price\n- adult\n- child"
    _save_workbook(input_dir / "inline-header.xlsx", workbook)

    _write_text(
        config_path,
        """
[files."inline-header.xlsx".sheets."Offers"]
parser_profile = "outline_sheet"
logical_id = "inline-header"
category = "sales"
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 0
    assert result.failure_count == 1
    assert result.errors[0].diagnostics[0].code == "inline_header_in_same_cell_unsupported"
    assert result.errors[0].diagnostics[0].row_numbers == (2,)


def test_xlsx_outline_sheet_keeps_plain_multiline_prose_as_paragraph(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    multiline_offer = (
        "\u0414\u0456\u0442\u0438 \u0434\u043e 3 \u0440\u043e\u043a\u0456\u0432\n"
        "\u0432\u043a\u043b\u044e\u0447\u043d\u043e \u0442\u0430 "
        "\u0456\u043c\u0435\u043d\u0438\u043d\u043d\u0438\u043a\u0438 "
        "\u0432\u0445\u043e\u0434\u044f\u0442\u044c "
        "\u0431\u0435\u0437\u043a\u043e\u0448\u0442\u043e\u0432\u043d\u043e"
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land"
    sheet["A2"] = multiline_offer
    _save_workbook(input_dir / "offers.xlsx", workbook)

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=None,
        source_formats=("xlsx",),
    )

    assert result.success_count == 1
    markdown = (output_dir / "markdown" / "offers-offers.md").read_text(encoding="utf-8")
    assert "###" not in markdown
    assert "\u0414\u0456\u0442\u0438 \u0434\u043e 3 \u0440\u043e\u043a\u0456\u0432" in markdown
    assert (
        "\u0432\u043a\u043b\u044e\u0447\u043d\u043e \u0442\u0430 "
        "\u0456\u043c\u0435\u043d\u0438\u043d\u043d\u0438\u043a\u0438 "
        "\u0432\u0445\u043e\u0434\u044f\u0442\u044c "
        "\u0431\u0435\u0437\u043a\u043e\u0448\u0442\u043e\u0432\u043d\u043e"
        in markdown
    )


def test_forced_header_rows_override_labelled_cells(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land"
    sheet["A2"] = "Price: 550 UAH"
    _save_workbook(input_dir / "offers.xlsx", workbook)

    _write_text(
        config_path,
        """
[files."offers.xlsx".sheets."Offers"]
forced_header_rows = [2]
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
        source_formats=("xlsx",),
    )

    assert result.success_count == 1
    markdown = (output_dir / "markdown" / "offers-offers.md").read_text(encoding="utf-8")
    assert "## Price: 550 UAH" in markdown
    assert "### Price" not in markdown


def test_xlsx_sheet_can_use_explicit_table_profile(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FAQ"
    sheet["A1"] = "Question"
    sheet["B1"] = "Answer"
    sheet["A2"] = "How to order?"
    sheet["B2"] = "Use Telegram."
    _save_workbook(input_dir / "faq.xlsx", workbook)

    _write_text(
        config_path,
        """
[files."faq.xlsx".sheets."FAQ"]
parser_profile = "qa_table"
question_column_name = "Question"
answer_column_name = "Answer"
logical_id = "faq"
category = "general"
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 1
    assert result.failure_count == 0
    markdown = (output_dir / "markdown" / "faq-faq.md").read_text(encoding="utf-8")
    assert markdown == "# Faq\n\n### How to order?\n\nUse Telegram.\n"


def test_config_can_disable_csv_parsing_and_select_sheet_indexes(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    _write_text(input_dir / "legacy.csv", "Question,Answer\nQ1,A1\n")

    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.merge_cells("A1:C1")
    first["A1"] = "First Sheet"
    first["A2"] = "Alpha details."
    second = workbook.create_sheet("Second")
    second.merge_cells("A1:C1")
    second["A1"] = "Second Sheet"
    second["A2"] = "Beta details."
    third = workbook.create_sheet("Third")
    third.merge_cells("A1:C1")
    third["A1"] = "Third Sheet"
    third["A2"] = "Gamma details."
    _save_workbook(input_dir / "bundle.xlsx", workbook)

    _write_text(
        config_path,
        """
[defaults]
source_formats = ["xlsx"]

[files."bundle.xlsx"]
sheet_indexes = [1, 3]
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 2
    assert result.failure_count == 0
    assert [entry.sheet_name for entry in result.entries] == ["First", "Third"]
    assert all(entry.source_format == "xlsx" for entry in result.entries)
    assert not (output_dir / "markdown" / "legacy.md").exists()


def test_sheet_indexes_use_visible_order_when_hidden_tabs_exist(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    hidden = workbook.active
    hidden.title = "Hidden"
    hidden.sheet_state = "hidden"
    visible = workbook.create_sheet("Visible")
    visible.merge_cells("A1:C1")
    visible["A1"] = "Visible Sheet"
    visible["A2"] = "Body."
    _save_workbook(input_dir / "bundle.xlsx", workbook)

    _write_text(
        config_path,
        """
[defaults]
source_formats = ["xlsx"]

[files."bundle.xlsx"]
sheet_indexes = [1]
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 1
    assert result.failure_count == 0
    assert [entry.sheet_name for entry in result.entries] == ["Visible"]
    assert result.entries[0].sheet_index == 1


def test_invalid_sheet_indexes_fail_instead_of_under_parsing(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first["A1"] = "First"
    second = workbook.create_sheet("Second")
    second["A1"] = "Second"
    _save_workbook(input_dir / "bundle.xlsx", workbook)

    _write_text(
        config_path,
        """
[defaults]
source_formats = ["xlsx"]

[files."bundle.xlsx"]
sheet_indexes = [0, 2, 999]
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 0
    assert result.failure_count == 1
    assert result.errors[0].error_type == "ValueError"
    assert "0, 999" in result.errors[0].message
    assert "Available visible sheet indexes: 1, 2" in result.errors[0].message


def test_cli_source_format_override_can_enable_csv_when_config_defaults_to_xlsx(
    tmp_path: Path,
) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    _write_text(input_dir / "faq.csv", "Question,Answer\nQ1,A1\n")
    _write_text(
        config_path,
        """
[defaults]
source_formats = ["xlsx"]

[files."faq.csv"]
question_column_name = "Question"
answer_column_name = "Answer"
""".strip(),
    )

    exit_code = main(
        [
            "--input-dir",
            str(input_dir),
            "--output-dir",
            str(output_dir),
            "--config",
            str(config_path),
            "--source-format",
            "csv",
        ]
    )

    assert exit_code == 0
    assert (output_dir / "markdown" / "faq.md").exists()


def test_xlsx_cyrillic_names_generate_distinct_transliterated_files(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"
    config_path = tmp_path / "parser_config.toml"

    workbook_name = (
        "\u041f\u043e\u0432\u0456\u0434\u043e\u043c\u043b\u0435\u043d\u043d\u044f 26.xlsx"
    )
    sheet_one = "\u041e\u041f \u0432\u0435\u0441\u043d\u0430-\u043b\u0456\u0442\u043e"
    sheet_two = "\u041a\u0435\u043c\u043f\u0456\u043d\u0433"

    workbook = Workbook()
    first = workbook.active
    first.title = sheet_one
    first.merge_cells("A1:C1")
    first["A1"] = "First title"
    first["A2"] = "First body."
    second = workbook.create_sheet(sheet_two)
    second.merge_cells("A1:C1")
    second["A1"] = "Second title"
    second["A2"] = "Second body."
    _save_workbook(input_dir / workbook_name, workbook)

    markdown_dir = output_dir / "markdown"
    markdown_dir.mkdir(parents=True, exist_ok=True)
    _write_text(markdown_dir / "26.md", "# stale\n")

    _write_text(
        config_path,
        f"""
[defaults]
source_formats = ["xlsx"]

[files."{workbook_name}"]
sheet_indexes = [1, 2]
""".strip(),
    )

    result = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=config_path,
    )

    assert result.success_count == 2
    assert result.failure_count == 0
    generated_names = sorted(path.name for path in markdown_dir.glob("*.md"))
    assert generated_names == [
        "povidomlennia-26-kempinh.md",
        "povidomlennia-26-op-vesna-lito.md",
    ]


def test_failed_rerun_removes_stale_markdown_after_new_ambiguity(tmp_path: Path) -> None:
    input_dir = tmp_path / "raw"
    output_dir = tmp_path / "processed"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land"
    sheet["A2"] = "Safe paragraph."
    _save_workbook(input_dir / "demo.xlsx", workbook)

    first = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=None,
        source_formats=("xlsx",),
    )
    assert first.success_count == 1
    assert (output_dir / "markdown" / "demo-offers.md").exists()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Offers"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Berry Land"
    sheet["A2"] = "Price\n- adult\n- child"
    _save_workbook(input_dir / "demo.xlsx", workbook)

    second = parse_knowledge_base(
        input_dir=input_dir,
        output_dir=output_dir,
        config_path=None,
        source_formats=("xlsx",),
    )

    assert second.success_count == 0
    assert second.failure_count == 1
    assert not (output_dir / "markdown" / "demo-offers.md").exists()


def test_infer_fallback_headers_does_not_guess_headers_for_two_column_pairs() -> None:
    rows = [
        ["Dates", "May 1"],
        ["Location", "Kyiv"],
        ["Schedule", "09:00"],
    ]

    inferred_headers, data_rows = infer_fallback_headers(rows)

    assert inferred_headers is None
    assert data_rows == rows


def test_infer_fallback_headers_accepts_larger_header_like_tables() -> None:
    rows = [
        ["Question", "Answer", "Section"],
        ["How to register?", "Fill in the form.", "General"],
        ["When does it open?", "At 09:00.", "General"],
    ]

    inferred_headers, data_rows = infer_fallback_headers(rows)

    assert inferred_headers == rows[0]
    assert data_rows == rows[1:]
