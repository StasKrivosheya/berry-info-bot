from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries

from app.services.knowledge_base.csv_ingest import PreparedTable
from app.services.knowledge_base.normalizer import normalize_cell_text
from app.services.knowledge_base.types import SourceOverride

COLUMN_RANGE_RE = re.compile(r"^(?P<start>[A-Za-z]+):(?P<end>[A-Za-z]+)$")
SINGLE_COLUMN_RE = re.compile(r"^(?P<column>[A-Za-z]+)$")


@dataclass(frozen=True, slots=True)
class SheetRegion:
    min_col: int
    max_col: int
    min_row: int | None = None
    max_row: int | None = None

    def contains(self, row_index: int, column_index: int) -> bool:
        if column_index < self.min_col or column_index > self.max_col:
            return False
        if self.min_row is not None and row_index < self.min_row:
            return False
        if self.max_row is not None and row_index > self.max_row:
            return False
        return True


@dataclass(frozen=True, slots=True)
class SheetCell:
    row_index: int
    column_index: int
    coordinate: str
    text: str
    is_bold: bool
    merged_span: int


@dataclass(frozen=True, slots=True)
class SheetRow:
    row_index: int
    cells: tuple[SheetCell, ...]
    row_height: float | None

    @property
    def primary_text(self) -> str:
        if not self.cells:
            return ""
        return self.cells[0].text

    @property
    def max_merged_span(self) -> int:
        if not self.cells:
            return 1
        return max(cell.merged_span for cell in self.cells)

    def with_cells(self, cells: tuple[SheetCell, ...]) -> SheetRow:
        return SheetRow(row_index=self.row_index, cells=cells, row_height=self.row_height)


@dataclass(frozen=True, slots=True)
class PreparedWorkbookSheet:
    workbook_file: str
    sheet_index: int
    sheet_name: str
    rows: tuple[SheetRow, ...]


def load_visible_workbook_sheets(workbook_path: Path) -> list[PreparedWorkbookSheet]:
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        sheets: list[PreparedWorkbookSheet] = []
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            if worksheet.sheet_state != "visible":
                continue
            sheets.append(_prepare_sheet(workbook_path.name, sheet_index, worksheet))
        return sheets
    finally:
        workbook.close()


def select_outline_rows(sheet: PreparedWorkbookSheet, override: SourceOverride) -> list[SheetRow]:
    content_regions = _parse_regions(override.content_ranges)
    ignored_regions = _parse_regions(override.ignore_columns)
    ignored_rows = set(override.ignore_rows)

    selected_rows: list[SheetRow] = []
    for row in sheet.rows:
        if row.row_index in ignored_rows:
            continue
        cells = tuple(
            cell
            for cell in row.cells
            if _cell_allowed(
                row_index=cell.row_index,
                column_index=cell.column_index,
                content_regions=content_regions,
                ignored_regions=ignored_regions,
            )
        )
        if cells:
            selected_rows.append(row.with_cells(cells))
    return selected_rows


def build_table_from_sheet(sheet: PreparedWorkbookSheet, override: SourceOverride) -> PreparedTable:
    selected_rows = select_outline_rows(sheet, override)
    if not selected_rows:
        return PreparedTable(rows=[], row_count=0, non_empty_cell_count=0)

    used_columns = sorted({cell.column_index for row in selected_rows for cell in row.cells})
    column_positions = {
        column_index: position for position, column_index in enumerate(used_columns)
    }
    rows: list[list[str]] = []
    non_empty_cell_count = 0
    for row in selected_rows:
        prepared_row = [""] * len(used_columns)
        for cell in row.cells:
            prepared_row[column_positions[cell.column_index]] = cell.text
            non_empty_cell_count += 1
        rows.append(prepared_row)

    return PreparedTable(
        rows=rows,
        row_count=len(rows),
        non_empty_cell_count=non_empty_cell_count,
    )


def _prepare_sheet(workbook_file: str, sheet_index: int, worksheet) -> PreparedWorkbookSheet:
    merged_spans = _build_merged_spans(worksheet)
    rows: list[SheetRow] = []
    for row_index in range(1, worksheet.max_row + 1):
        cells: list[SheetCell] = []
        for column_index in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            text = _normalize_cell_value(cell.value)
            if not text:
                continue
            cells.append(
                SheetCell(
                    row_index=row_index,
                    column_index=column_index,
                    coordinate=cell.coordinate,
                    text=text,
                    is_bold=bool(getattr(cell.font, "bold", False)),
                    merged_span=merged_spans.get(cell.coordinate, 1),
                )
            )
        if not cells:
            continue
        row_height = worksheet.row_dimensions[row_index].height
        rows.append(
            SheetRow(
                row_index=row_index,
                cells=tuple(cells),
                row_height=float(row_height) if row_height is not None else None,
            )
        )
    return PreparedWorkbookSheet(
        workbook_file=workbook_file,
        sheet_index=sheet_index,
        sheet_name=str(worksheet.title),
        rows=tuple(rows),
    )


def _build_merged_spans(worksheet) -> dict[str, int]:
    spans: dict[str, int] = {}
    for merged_range in worksheet.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        top_left = worksheet.cell(row=min_row, column=min_col).coordinate
        if min_row == max_row:
            spans[top_left] = max(1, max_col - min_col + 1)
    return spans


def _normalize_cell_value(raw_value: object) -> str:
    if raw_value is None:
        return ""
    return normalize_cell_text(str(raw_value))


def _parse_regions(tokens: tuple[str, ...]) -> tuple[SheetRegion, ...]:
    return tuple(_parse_region(token) for token in tokens)


def _parse_region(token: str) -> SheetRegion:
    normalized = normalize_cell_text(token)
    single_column_match = SINGLE_COLUMN_RE.match(normalized)
    if single_column_match:
        column_index = column_index_from_string(single_column_match.group("column"))
        return SheetRegion(min_col=column_index, max_col=column_index)

    column_range_match = COLUMN_RANGE_RE.match(normalized)
    if column_range_match:
        min_col = column_index_from_string(column_range_match.group("start"))
        max_col = column_index_from_string(column_range_match.group("end"))
        return SheetRegion(min_col=min_col, max_col=max_col)

    min_col, min_row, max_col, max_row = range_boundaries(normalized)
    return SheetRegion(
        min_col=min_col,
        max_col=max_col,
        min_row=min_row,
        max_row=max_row,
    )


def _cell_allowed(
    *,
    row_index: int,
    column_index: int,
    content_regions: tuple[SheetRegion, ...],
    ignored_regions: tuple[SheetRegion, ...],
) -> bool:
    if ignored_regions and any(
        region.contains(row_index, column_index) for region in ignored_regions
    ):
        return False
    if not content_regions:
        return True
    return any(region.contains(row_index, column_index) for region in content_regions)
